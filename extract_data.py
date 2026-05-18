#!/usr/bin/env python3
"""Extract all process sheets into a structured database."""
import os, re, sqlite3, json
from pathlib import Path
from docx import Document

SRC = os.path.expanduser("~/Desktop/生产工艺单")
DB = os.path.join(os.path.dirname(__file__), "data", "process_sheets.db")

def extract_sheet(filepath):
    """Extract key fields from a single process sheet .docx"""
    try:
        doc = Document(filepath)
        tables = doc.tables
        if len(tables) < 3:
            return None
        
        data = {
            'file': os.path.basename(filepath),
            'dirname': os.path.basename(os.path.dirname(filepath)),
            'filepath': filepath,
        }
        
        # Table 0: Customer, Product, Qty, Dimensions
        t0 = tables[0]
        rows0 = {c: [] for c in range(len(t0.rows[0].cells))}
        for row in t0.rows:
            for ci, cell in enumerate(row.cells):
                rows0.setdefault(ci, []).append(cell.text.strip())
        
        # Flatten and find key fields
        all_text_0 = ' '.join([' '.join(rows0[c]) for c in rows0])
        
        # Parse header line for batch no, date
        for p in doc.paragraphs:
            t = p.text.strip()
            if '批次号' in t:
                m = re.search(r'批次号[：:]\s*(\S+)', t)
                if m: data['batch_no'] = m.group(1)
            if '下单时间' in t:
                m = re.search(r'下单时间[：:]\s*(\S+)', t)
                if m: data['order_date'] = m.group(1)
        
        # Extract from table 0
        t0_text = all_text_0
        # Client
        # 客户名提取：用通用字段标签作 boundary，不硬编码任何真实公司名
        m = re.search(r'客户名称\s*(\S.{0,30}?)(?=\s{2,}|产品名称|订单号|批号|交货|$)', t0_text)
        if m: data['client'] = m.group(1).strip()
        
        # Product name
        m = re.search(r'产品名称\s*(\S.*?)(?:交货|成品数量|订单数量|规格)', t0_text)
        if m: data['product'] = m.group(1).strip()
        
        # Quantity
        m = re.search(r'订单数量\s*(\d+(?:\.\d+)?)\s*只', t0_text)
        if m: data['order_qty'] = int(float(m.group(1)))
        m = re.search(r'成品数量\s*(\d+(?:\.\d+)?)\s*只', t0_text)
        if m: data['finished_qty'] = int(float(m.group(1)))
        
        # Dimensions (L*W*H)
        m = re.search(r'规格\s*(\d+)\*(\d+)\*(\d+)', t0_text)
        if m:
            data['box_l'] = int(m.group(1))
            data['box_w'] = int(m.group(2))
            data['box_h'] = int(m.group(3))
        
        # Notes
        notes = []
        for c in rows0:
            for txt in rows0[c]:
                if '备注' in txt:
                    notes.append(txt)
        data['notes'] = '; '.join(notes)
        
        # Table 1: Paper specs (面纸)
        t1 = tables[1]
        t1_text = ' '.join([' '.join([cell.text.strip() for cell in row.cells]) for row in t1.rows])
        
        m = re.search(r'面纸尺寸[，,]*mm\s*(\d+)\*(\d+)', t1_text)
        if m:
            data['paper_w'] = int(m.group(1))
            data['paper_h'] = int(m.group(2))
        
        m = re.search(r'面纸规格型号\s*(\S.*?)(?:覆膜|亮膜|哑膜|上光)', t1_text)
        if m: data['paper_spec'] = m.group(1).strip()
        
        m = re.search(r'印刷数量[，,]*张\s*(\d+)', t1_text)
        if m: data['print_qty'] = int(m.group(1))
        
        # Lamination
        if '亮膜' in t1_text and '是' in t1_text.split('亮膜')[1][:5] if '亮膜' in t1_text else False:
            data['lamination'] = '亮膜'
        elif '哑膜' in t1_text:
            data['lamination'] = '哑膜' if '是' in t1_text else None
        
        # Find lamination more carefully
        m = re.search(r'覆膜.*?(亮膜|哑膜|上光)', t1_text)
        data['has_lamination'] = '亮膜' in t1_text or '哑膜' in t1_text
        
        # Table 2: Corrugated specs (瓦楞纸板)
        t2 = tables[2]
        t2_text = ' '.join([' '.join([cell.text.strip() for cell in row.cells]) for row in t2.rows])
        
        m = re.search(r'瓦楞纸板尺寸[，,]*mm\s*(\d+)\*(\d+)', t2_text)
        if m:
            data['board_w'] = int(m.group(1))
            data['board_h'] = int(m.group(2))
        
        m = re.search(r'瓦型\s*(EB|BC|单C瓦|单B瓦|BE|EC|EE|BB)', t2_text)
        if m: data['flute_type'] = m.group(1)
        
        m = re.search(r'纸板材质\s*(\S.*?)(?:垫片|格档|箱钉|压痕|正压|反压|——)', t2_text)
        if m: data['board_material'] = m.group(1).strip()
        
        # Padding/Grid
        m = re.search(r'垫片材质\s*(\S.*?)(?:垫片|尺寸|数量)', t2_text)
        if m: data['padding_material'] = m.group(1).strip()
        m = re.search(r'垫片[，,]尺寸[，,]mm\s*(\d+)\*(\d+)', t2_text)
        if m: data['padding_size'] = f"{m.group(1)}*{m.group(2)}"
        m = re.search(r'格档.*?(BC瓦|单B瓦|\S+瓦).*?(\d+g)', t2_text)
        if m: data['grid_material'] = m.group(0)
        
        # Detect print style
        if '双拼' in data.get('notes', ''):
            data['print_style'] = '双拼'
        elif '两页' in data.get('notes', ''):
            data['print_style'] = '两页成型'
        elif '单页' in data.get('notes', ''):
            data['print_style'] = '单页成型'
        
        # Detect if it's a quote (报价)
        data['is_quote'] = '报价' in data.get('file', '') or '报价' in data.get('dirname', '')
        
        return data
    except Exception as e:
        return {'file': os.path.basename(filepath), 'error': str(e)}


def main():
    conn = sqlite3.connect(DB)
    
    # Create table
    conn.execute('''
        CREATE TABLE IF NOT EXISTS process_sheets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            file TEXT, dirname TEXT, filepath TEXT,
            batch_no TEXT, order_date TEXT, client TEXT, product TEXT,
            order_qty INTEGER, finished_qty INTEGER,
            box_l INTEGER, box_w INTEGER, box_h INTEGER,
            paper_w INTEGER, paper_h INTEGER, paper_spec TEXT,
            print_qty INTEGER, lamination TEXT, has_lamination BOOLEAN,
            board_w INTEGER, board_h INTEGER, flute_type TEXT, board_material TEXT,
            padding_material TEXT, padding_size TEXT, grid_material TEXT,
            print_style TEXT, notes TEXT, is_quote BOOLEAN,
            error TEXT
        )
    ''')
    
    # Walk all directories
    total, success, failed = 0, 0, 0
    rows = []
    
    for root, dirs, files in os.walk(SRC):
        for f in files:
            if f.startswith('.') or not f.endswith('.docx'):
                continue
            total += 1
            fp = os.path.join(root, f)
            data = extract_sheet(fp)
            if data:
                if 'error' in data:
                    failed += 1
                else:
                    success += 1
                rows.append(data)
            
            if total % 500 == 0:
                print(f"Processed {total} files... ({success} ok, {failed} err)")
    
    # Batch insert
    cols = ['file','dirname','filepath','batch_no','order_date','client','product',
            'order_qty','finished_qty','box_l','box_w','box_h',
            'paper_w','paper_h','paper_spec','print_qty','lamination','has_lamination',
            'board_w','board_h','flute_type','board_material',
            'padding_material','padding_size','grid_material',
            'print_style','notes','is_quote','error']
    
    for r in rows:
        vals = [r.get(c) for c in cols]
        placeholders = ','.join(['?']*len(cols))
        conn.execute(f'INSERT INTO process_sheets ({",".join(cols)}) VALUES ({placeholders})', vals)
    
    conn.commit()
    
    # Stats
    print(f"\n=== Extraction Complete ===")
    print(f"Total files scanned: {total}")
    print(f"Successfully parsed: {success} ({success*100//max(total,1)}%)")
    print(f"Failed: {failed}")
    
    # Quick stats
    for col in ['client', 'flute_type', 'print_style', 'lamination']:
        cur = conn.execute(f"SELECT {col}, COUNT(*) FROM process_sheets WHERE error IS NULL AND {col} IS NOT NULL GROUP BY {col} ORDER BY COUNT(*) DESC LIMIT 10")
        print(f"\nTop {col}:")
        for row in cur.fetchall():
            print(f"  {row[0]}: {row[1]}")
    
    conn.close()
    print(f"\nDatabase saved to: {DB}")

if __name__ == '__main__':
    main()
